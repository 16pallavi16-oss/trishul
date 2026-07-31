from pathlib import Path
from typing import List, Dict, Any
import csv

from unstructured.partition.auto import partition
from openpyxl import load_workbook

MAX_CHUNK_CHARS = 1000

SUPPORTED_UNSTRUCTURED_EXTS = {".pdf", ".docx", ".pptx"}
SUPPORTED_EXTS = SUPPORTED_UNSTRUCTURED_EXTS | {".xlsx", ".csv"}

OCR_MIN_CHARS_THRESHOLD = 20

def _make_chunk(text: str, source: str, chunk_index: int, **extra: Any) -> Dict[str, Any]:
    chunk = {
        "text": text.strip(),
        "source": source,
        "chunk_index": chunk_index,
    }
    chunk.update(extra)
    return chunk


def _split_long_text(text: str, max_chars: int) -> List[str]:
    if len(text) <= max_chars:
        return [text]

    pieces = []
    current = ""
    for line in text.split("\n"):
        if len(current) + len(line) + 1 > max_chars and current:
            pieces.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        pieces.append(current)
    return pieces


def _elements_to_chunks(elements, path: Path, ocr_used: bool = False) -> List[Dict[str, Any]]:
    chunks = []
    idx = 0
    for element in elements:
        text = str(element).strip()
        if not text:
            continue

        category = getattr(element, "category", None)
        page_number = None
        if getattr(element, "metadata", None) is not None:
            page_number = getattr(element.metadata, "page_number", None)

        for piece in _split_long_text(text, MAX_CHUNK_CHARS):
            chunks.append(
                _make_chunk(
                    piece,
                    source=str(path),
                    chunk_index=idx,
                    file_type=path.suffix.lower(),
                    category=category,
                    page_number=page_number,
                    ocr_used=ocr_used,
                )
            )
            idx += 1

    return chunks


def ingest_unstructured_file(path: Path) -> List[Dict[str, Any]]:
    elements = partition(filename=str(path))
    total_chars = sum(len(str(el).strip()) for el in elements)

    if path.suffix.lower() == ".pdf" and total_chars < OCR_MIN_CHARS_THRESHOLD:
        print(f"[ocr] {path.name}: little/no text found ({total_chars} chars) -- retrying with Tesseract OCR")
        try:
            elements = partition(filename=str(path), strategy="ocr_only")
            return _elements_to_chunks(elements, path, ocr_used=True)
        except Exception as e:
            print(f"[ocr-failed] {path.name}: OCR retry failed ({e}); keeping original (empty) extraction")

    return _elements_to_chunks(elements, path, ocr_used=False)


def ingest_xlsx_file(path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(filename=str(path), data_only=True)

    chunks = []
    idx = 0
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        headers = None

        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number == 1:
                headers = [str(c) if c is not None else "" for c in row]
                continue

            if all(cell is None for cell in row):
                continue

            if headers:
                row_text = ", ".join(
                    f"{headers[i]}: {row[i]}"
                    for i in range(len(row))
                    if row[i] is not None
                )
            else:
                row_text = ", ".join(str(cell) for cell in row if cell is not None)

            chunks.append(
                _make_chunk(
                    row_text,
                    source=str(path),
                    chunk_index=idx,
                    file_type=".xlsx",
                    sheet_name=sheet_name,
                    row_number=row_number,
                )
            )
            idx += 1

    return chunks


def ingest_csv_file(path: Path) -> List[Dict[str, Any]]:
    chunks = []
    idx = 0

    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row_number, row in enumerate(reader, start=1):
            row_text = ", ".join(f"{k}: {v}" for k, v in row.items() if v)
            if not row_text.strip():
                continue

            chunks.append(
                _make_chunk(
                    row_text,
                    source=str(path),
                    chunk_index=idx,
                    file_type=".csv",
                    row_number=row_number,
                )
            )
            idx += 1

    return chunks

def ingest_file(path: Path) -> List[Dict[str, Any]]:
    ext = path.suffix.lower()

    if ext in SUPPORTED_UNSTRUCTURED_EXTS:
        return ingest_unstructured_file(path)
    elif ext == ".xlsx":
        return ingest_xlsx_file(path)
    elif ext == ".csv":
        return ingest_csv_file(path)
    else:
        return []


def ingest_folder(folder_path: str, recursive: bool = True) -> List[Dict[str, Any]]:
    root = Path(folder_path)
    if not root.exists():
        raise FileNotFoundError(f"Folder not found: {folder_path}")

    pattern = "**/*" if recursive else "*"
    all_chunks: List[Dict[str, Any]] = []

    for file_path in sorted(root.glob(pattern)):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SUPPORTED_EXTS:
            continue

        try:
            file_chunks = ingest_file(file_path)
            all_chunks.extend(file_chunks)
            print(f"[ok] {file_path.name}: {len(file_chunks)} chunks")
        except Exception as e:
            print(f"[skip] {file_path.name}: failed to ingest ({e})")

    return all_chunks


if __name__ == "__main__":
    import sys
    import json

    if len(sys.argv) != 2:
        print("Usage: python ingest.py <folder_path>")
        sys.exit(1)

    result = ingest_folder(sys.argv[1])
    print(f"\nTotal chunks: {len(result)}")
    print(json.dumps(result[:3], indent=2))