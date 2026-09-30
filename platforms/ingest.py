from pathlib import Path
from typing import List, Dict, Any
import csv
import os

from unstructured.partition.auto import partition
from openpyxl import load_workbook
import pypdfium2 as pdfium
import pytesseract
from PIL import Image

MAX_CHUNK_CHARS = 1000
CSV_ROWS_PER_CHUNK = 100

SUPPORTED_UNSTRUCTURED_EXTS = {".pdf", ".docx", ".pptx"}
SUPPORTED_EXTS = SUPPORTED_UNSTRUCTURED_EXTS | {".xlsx", ".csv"}

OCR_MIN_CHARS_THRESHOLD = 20
TESSERACT_MIN_CHARS_THRESHOLD = 20  # below this, escalate a page from Tesseract to docTR

def _make_chunk(text: str, source: str, chunk_index: int, **extra: Any) -> Dict[str, Any]:
    chunk = {
        "text": text.strip(),
        "source": source,
        "chunk_index": chunk_index,
    }
    chunk.update(extra)
    return chunk


def _split_long_text(text: str, max_chars: int) -> List[str]:
    if len(text) <= max_chars:
        return [text]

    pieces = []
    current = ""
    for line in text.split("\n"):
        if len(current) + len(line) + 1 > max_chars and current:
            pieces.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        pieces.append(current)
    return pieces


def _elements_to_chunks(elements, path: Path, ocr_used: bool = False) -> List[Dict[str, Any]]:
    chunks = []
    idx = 0
    for element in elements:
        text = str(element).strip()
        if not text:
            continue

        category = getattr(element, "category", None)
        page_number = None
        if getattr(element, "metadata", None) is not None:
            page_number = getattr(element.metadata, "page_number", None)

        for piece in _split_long_text(text, MAX_CHUNK_CHARS):
            chunks.append(
                _make_chunk(
                    piece,
                    source=str(path),
                    chunk_index=idx,
                    file_type=path.suffix.lower(),
                    category=category,
                    page_number=page_number,
                    ocr_used=ocr_used,
                )
            )
            idx += 1

    return chunks


_doctr_model = None  # lazy singleton -- the layout-aware model only loads if a page actually needs it

def _get_doctr_model():
    """Load docTR's OCR model on first use only (it's slow to load and most pages won't need it)."""
    global _doctr_model
    if _doctr_model is None:
        os.environ.setdefault("USE_TORCH", "1")
        from doctr.models import ocr_predictor
        print("[doctr] loading layout-aware OCR model (first use only, downloads weights if needed)...")
        _doctr_model = ocr_predictor(pretrained=True)
    return _doctr_model


def _render_pdf_pages(path: Path, dpi: int = 300) -> List[Image.Image]:
    """Rasterize every page of a PDF to a PIL image via pypdfium2 -- no Poppler involved."""
    pdf = pdfium.PdfDocument(str(path))
    images = []
    for page in pdf:
        bitmap = page.render(scale=dpi / 72)
        images.append(bitmap.to_pil())
    return images


def _ocr_pdf_file(path: Path) -> List[Dict[str, Any]]:
    """
    Two-tier OCR for scanned/no-text PDFs:
      1. Fast pass -- Tesseract on each rasterized page.
      2. Layout-aware fallback -- docTR, only for pages where Tesseract found too little text.
    Rendering (pypdfium2) is shared by both tiers and never touches Poppler.
    """
    images = _render_pdf_pages(path)
    chunks: List[Dict[str, Any]] = []
    idx = 0

    for page_number, image in enumerate(images, start=1):
        text = pytesseract.image_to_string(image).strip()

        if len(text) >= TESSERACT_MIN_CHARS_THRESHOLD:
            for piece in _split_long_text(text, MAX_CHUNK_CHARS):
                chunks.append(
                    _make_chunk(
                        piece, source=str(path), chunk_index=idx,
                        file_type=".pdf", category="OCR-Text",
                        page_number=page_number, ocr_used="tesseract",
                    )
                )
                idx += 1
            continue

        print(f"[ocr-fallback] {path.name} p{page_number}: Tesseract found little text -- trying docTR")
        import numpy as np
        model = _get_doctr_model()
        result = model([np.array(image)])
        doctr_page = result.pages[0]

        for block in doctr_page.blocks:
            for line in block.lines:
                line_text = " ".join(word.value for word in line.words).strip()
                if not line_text:
                    continue
                (x0, y0), (x1, y1) = line.geometry
                chunks.append(
                    _make_chunk(
                        line_text, source=str(path), chunk_index=idx,
                        file_type=".pdf", category="OCR-Text",
                        page_number=page_number, bbox=[x0, y0, x1, y1],
                        ocr_used="doctr",
                    )
                )
                idx += 1

    return chunks


def ingest_unstructured_file(path: Path) -> List[Dict[str, Any]]:
    partition_kwargs = {"filename": str(path)}
    if path.suffix.lower() == ".pdf":
        partition_kwargs["strategy"] = "fast"  # pdfminer text extraction only -- no Poppler needed

    elements = partition(**partition_kwargs)
    total_chars = sum(len(str(el).strip()) for el in elements)

    if path.suffix.lower() == ".pdf" and total_chars < OCR_MIN_CHARS_THRESHOLD:
        print(f"[ocr] {path.name}: little/no text found ({total_chars} chars) -- rasterizing + Tesseract/docTR OCR")
        try:
            return _ocr_pdf_file(path)
        except Exception as e:
            print(f"[ocr-failed] {path.name}: OCR failed ({e}); keeping original (empty) extraction")

    return _elements_to_chunks(elements, path, ocr_used=False)


def ingest_xlsx_file(path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(filename=str(path), data_only=True)

    chunks = []
    idx = 0
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        headers = None

        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number == 1:
                headers = [str(c) if c is not None else "" for c in row]
                continue

            if all(cell is None for cell in row):
                continue

            if headers:
                row_text = ", ".join(
                    f"{headers[i]}: {row[i]}"
                    for i in range(len(row))
                    if row[i] is not None
                )
            else:
                row_text = ", ".join(str(cell) for cell in row if cell is not None)

            chunks.append(
                _make_chunk(
                    row_text,
                    source=str(path),
                    chunk_index=idx,
                    file_type=".xlsx",
                    sheet_name=sheet_name,
                    row_number=row_number,
                )
            )
            idx += 1

    return chunks


def ingest_csv_file(path: Path) -> List[Dict[str, Any]]:
    chunks = []
    idx = 0
    rows = []
    first_row_number = None

    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row_number, row in enumerate(reader, start=1):
            row_text = ", ".join(f"{k}: {v}" for k, v in row.items() if v)
            if not row_text.strip():
                continue

            if first_row_number is None:
                first_row_number = row_number
            rows.append(row_text)

            if len(rows) < CSV_ROWS_PER_CHUNK:
                continue

            chunks.append(_make_chunk(
                "\n".join(rows),
                source=str(path),
                chunk_index=idx,
                file_type=".csv",
                row_start=first_row_number,
                row_end=row_number,
            ))
            idx += 1
            rows = []
            first_row_number = None

    if rows:
        chunks.append(_make_chunk(
            "\n".join(rows),
            source=str(path),
            chunk_index=idx,
            file_type=".csv",
            row_start=first_row_number,
            row_end=first_row_number + len(rows) - 1,
        ))

    return chunks

def ingest_file(path: Path) -> List[Dict[str, Any]]:
    ext = path.suffix.lower()

    if ext in SUPPORTED_UNSTRUCTURED_EXTS:
        return ingest_unstructured_file(path)
    elif ext == ".xlsx":
        return ingest_xlsx_file(path)
    elif ext == ".csv":
        return ingest_csv_file(path)
    else:
        return []


def ingest_folder(folder_path: str, recursive: bool = True) -> List[Dict[str, Any]]:
    root = Path(folder_path)
    if not root.exists():
        raise FileNotFoundError(f"Folder not found: {folder_path}")

    pattern = "**/*" if recursive else "*"
    all_chunks: List[Dict[str, Any]] = []

    for file_path in sorted(root.glob(pattern)):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SUPPORTED_EXTS:
            continue

        try:
            file_chunks = ingest_file(file_path)
            all_chunks.extend(file_chunks)
            print(f"[ok] {file_path.name}: {len(file_chunks)} chunks")
        except Exception as e:
            print(f"[skip] {file_path.name}: failed to ingest ({e})")

    return all_chunks


if __name__ == "__main__":
    import sys
    sys.modules.setdefault("platforms.ingest", sys.modules["__main__"])
    from . import db

    if len(sys.argv) != 2:
        print("Usage: python -m platforms.ingest <folder_path>")
        sys.exit(1)

    db.init_db()
    db.persist_folder(sys.argv[1])